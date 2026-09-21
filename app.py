"""Hugoland na nuvem: painel web + API + coletor contínuo."""
import hashlib
import hmac
import io
import json
import logging
import os
import re
import time
from datetime import datetime, timedelta, timezone
from functools import wraps
from zoneinfo import ZoneInfo

from flask import Flask, Response, jsonify, redirect, request, send_from_directory, session
from openpyxl import Workbook

import collector
from store import store

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(name)s %(levelname)s %(message)s")
BASE = os.path.dirname(os.path.abspath(__file__))
TZ = ZoneInfo(os.environ.get("TZ_APP", "America/Sao_Paulo"))
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")

app = Flask(__name__, static_folder=None)
app.secret_key = os.environ.get("SECRET_KEY") or hashlib.sha256(("hugoland|" + APP_PASSWORD + "|" + os.environ.get("DATABASE_URL", "")).encode()).hexdigest()
app.config.update(PERMANENT_SESSION_LIFETIME=timedelta(days=30), SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SECURE=os.environ.get("RAILWAY_ENVIRONMENT") is not None,
                  MAX_CONTENT_LENGTH=200 * 1024 * 1024)

store.init()
collector.start()

_last_sync = [0.0]


def fresh():
    if time.time() - _last_sync[0] > 3:
        store.sync()
        _last_sync[0] = time.time()


# ---------------- login ----------------
LOGIN_HTML = """<!DOCTYPE html><html lang="pt-BR"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Hugoland · Entrar</title><link rel="icon" href="/static/icon.png"><style>
body{margin:0;min-height:100vh;display:grid;place-items:center;background:radial-gradient(900px 500px at 20% -10%,#2B1F55,transparent 60%),#120D22;color:#F4F0FF;font-family:Rubik,system-ui,sans-serif}
form{width:min(360px,90vw);background:#211A3B;border:1px solid rgba(255,255,255,.1);border-radius:22px;padding:28px;text-align:center;box-shadow:0 20px 50px rgba(0,0,0,.4)}
img{width:96px;height:96px;border-radius:22px}h1{margin:12px 0 4px;font-size:1.6rem}h1 span{color:#F5C542}p{color:#A59CC6;margin:0 0 18px;font-size:.9rem}
input{width:100%;box-sizing:border-box;padding:13px;border-radius:12px;border:1px solid rgba(255,255,255,.15);background:#1A1330;color:#fff;font-size:1rem;margin-bottom:12px}
button{width:100%;padding:13px;border:0;border-radius:12px;background:linear-gradient(180deg,#FFE38A,#F5C542);color:#1A1330;font-weight:700;font-size:1rem;cursor:pointer}
.err{color:#FF4D6D;margin-bottom:10px;font-size:.9rem}</style></head><body>
<form method="post" action="/login"><img src="/static/icon.png" alt=""><h1>Hugo<span>land</span></h1><p>Monitor Adventures Beyond Wonderland</p>
__ERR__<input type="password" name="password" placeholder="Senha" autofocus autocomplete="current-password" required><button>Entrar</button></form></body></html>"""


def logged():
    return session.get("ok") is True


def need_login(f):
    @wraps(f)
    def w(*a, **k):
        if not APP_PASSWORD:
            return Response("Defina a variável APP_PASSWORD no Railway para liberar o acesso.", 503)
        if not logged():
            if request.path.startswith("/api/"):
                return jsonify(error="login necessário"), 401
            return redirect("/login")
        return f(*a, **k)
    return w


@app.get("/login")
def login_page():
    return LOGIN_HTML.replace("__ERR__", "")


@app.post("/login")
def login():
    if APP_PASSWORD and hmac.compare_digest(request.form.get("password", ""), APP_PASSWORD):
        session.permanent = True
        session["ok"] = True
        return redirect("/")
    time.sleep(1.5)
    return LOGIN_HTML.replace("__ERR__", '<div class="err">Senha incorreta</div>'), 401


@app.get("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.get("/health")
def health():
    return jsonify(ok=True, rodadas=len(store.times), coletor=collector.status.get("ok"))


# ---------------- páginas ----------------
@app.get("/")
@need_login
def index():
    return send_from_directory(os.path.join(BASE, "static"), "index.html")


@app.get("/static/<path:p>")
def static_files(p):
    return send_from_directory(os.path.join(BASE, "static"), p)


# ---------------- API ----------------
@app.get("/api/summary")
@need_login
def api_summary():
    fresh()
    s = store.summary()
    st = dict(collector.status)
    return jsonify(**s, collector=st, now=time.time() * 1000)


@app.get("/api/stream")
@need_login
def api_stream():
    """Avisa o navegador na hora em que uma rodada nova é gravada."""
    def gen():
        last = store.event_id
        yield "retry: 3000\n\n"
        while True:
            cur = store.wait_change(last, 20)
            if cur != last:
                last = cur
                yield f"event: change\ndata: {cur}\n\n"
            else:
                yield ": ping\n\n"
    return Response(gen(), mimetype="text/event-stream",
                    headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no", "Connection": "keep-alive"})


@app.get("/api/history")
@need_login
def api_history():
    fresh()
    off = max(0, int(request.args.get("offset", 0)))
    lim = min(500, max(1, int(request.args.get("limit", 100))))
    return jsonify(store.history(off, lim))


@app.post("/api/monitors")
@need_login
def api_monitor_save():
    d = request.get_json(force=True) or {}
    seg_ids = {s["id"] for s in store.segments}
    vals = [v for v in d.get("values", []) if v in seg_ids]
    if not vals:
        return jsonify(error="escolha pelo menos um resultado"), 400
    m = {"id": d.get("id"), "name": (d.get("name") or "").strip()[:80] or "Combinação", "values": vals,
         "mode": "seq" if d.get("mode") == "seq" else "any", "alert": max(0, int(d.get("alert") or 0))}
    return jsonify(ok=True, id=store.upsert_monitor(m))


@app.delete("/api/monitors/<mid>")
@need_login
def api_monitor_delete(mid):
    store.delete_monitor(mid)
    return jsonify(ok=True)


@app.put("/api/segments")
@need_login
def api_segments():
    segs = request.get_json(force=True) or []
    clean = []
    for s in segs:
        if not s.get("id"):
            continue
        color = s.get("color") if re.fullmatch(r"#[0-9a-fA-F]{6}", str(s.get("color", ""))) else "#8FA8FF"
        clean.append({"id": str(s["id"])[:20], "label": (str(s.get("label", "")).strip() or "?")[:24],
                      "color": color, "aliases": str(s.get("aliases", ""))[:200]})
    if not clean:
        return jsonify(error="lista vazia"), 400
    store.save_segments(clean)
    return jsonify(ok=True)


@app.post("/api/rounds")
@need_login
def api_manual():
    d = request.get_json(force=True) or {}
    code = store.code_for_seg(d.get("seg", ""))
    if not code:
        return jsonify(error="resultado inválido"), 400
    now = time.time()
    store.insert_rounds([(f"manual-{int(now * 1000)}", now, code)], src="manual")
    return jsonify(ok=True)


@app.delete("/api/rounds/last-manual")
@need_login
def api_undo():
    return jsonify(ok=store.delete_last_manual())


@app.delete("/api/rounds/<rid>")
@need_login
def api_delete_round(rid):
    return jsonify(ok=store.delete_round(rid))


@app.delete("/api/rounds")
@need_login
def api_clear():
    if (request.get_json(silent=True) or {}).get("confirm") != "APAGAR":
        return jsonify(error="confirmação ausente"), 400
    store.clear_rounds()
    return jsonify(ok=True)


# ---------------- exportação ----------------
def period_start():
    days = request.args.get("days", "all")
    if days == "all":
        return None, "tudo"
    d = max(1, int(days))
    return time.time() - d * 86400, f"últimos {d} dias"


def local(ts):
    return datetime.fromtimestamp(ts, TZ)


def seg_label(code):
    s = store.seg_for_code(code)
    return s["label"] if s else code


def export_rows(since):
    import numpy as np
    with store.lock:
        times, codes, srcs = store.times, store.codes, store.srcs
        start = int(np.searchsorted(times, since)) if since else 0
        idx = np.arange(start, len(times))
        seqs = [store.running(m, idx) for m in store.monitors] if len(idx) else []
        code_list = list(store.code_list)
    for k, i in enumerate(idx):
        yield int(i), float(times[i]), code_list[codes[i]], int(srcs[i]), [int(s[k]) for s in seqs]


def mode_txt(m):
    return "Zera quando sai a sequência" if m["mode"] == "seq" else "Zera quando sai qualquer um"


@app.get("/api/export.xlsx")
@need_login
def api_xlsx():
    fresh()
    since, label = period_start()
    wb = Workbook(write_only=True)
    segmap = {s["id"]: s["label"] for s in store.segments}
    ws = wb.create_sheet("Combinações")
    ws.append(["Combinação", "Resultados", "Regra", "Sequência atual", "Recorde", "Vezes que saiu", "Média entre saídas", "Última saída", "Alerta"])
    for m in store.monitors:
        st = store.stats(m)
        ws.append([m["name"], (" → " if m["mode"] == "seq" else ", ").join(segmap.get(v, "?") for v in m["values"]), mode_txt(m),
                   st["current"], st["record"], st["hits"], round(st["avg"], 2) if st["avg"] is not None else None,
                   local(st["lastT"] / 1000).strftime("%d/%m/%Y %H:%M:%S") if st["lastT"] else "nunca", m["alert"] or None])
    ws = wb.create_sheet("Rodadas")
    ws.append(["Nº", "Data", "Hora", "Resultado", "Código da fonte", "Origem"] + [f"Seq. {m['name']}" for m in store.monitors])
    for i, t, code, src, seqs in export_rows(since):
        dt = local(t)
        ws.append([i + 1, dt.strftime("%d/%m/%Y"), dt.strftime("%H:%M:%S"), seg_label(code), code, "Manual" if src else "Ao vivo", *seqs])
    ws = wb.create_sheet("Resumo")
    ws.append(["Resultado", "Vezes (histórico todo)", "% das rodadas", "Rodadas sem sair"])
    summ = store.summary(last=0)
    n = summ["n"] or 1
    for d in summ["dist"]:
        ws.append([segmap.get(d["seg"], "?"), d["count"], round(d["count"] / n * 100, 2), d["since"]])
    buf = io.BytesIO()
    wb.save(buf)
    name = f"hugoland_{datetime.now(TZ).strftime('%Y-%m-%d_%H%M')}.xlsx"
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/export.json")
@need_login
def api_json():
    fresh()
    since, label = period_start()
    segmap = {s["id"]: s["label"] for s in store.segments}

    def gen():
        head = {"app": "Hugoland", "versao": 3, "exportado_em": datetime.now(TZ).isoformat(), "periodo": label,
                "combinacoes": [{"nome": m["name"], "resultados": [segmap.get(v, "?") for v in m["values"]],
                                 "regra": "sequencia" if m["mode"] == "seq" else "qualquer_um", "alerta": m["alert"],
                                 **{k: v for k, v in store.stats(m).items() if k in ("current", "record", "hits", "avg")}}
                                for m in store.monitors]}
        yield json.dumps(head, ensure_ascii=False)[:-1] + ', "rodadas": ['
        names = [m["name"] for m in store.monitors]
        first = True
        for i, t, code, src, seqs in export_rows(since):
            row = {"n": i + 1, "data_hora": local(t).isoformat(), "resultado": seg_label(code), "codigo": code,
                   "origem": "manual" if src else "ao_vivo", "sequencias": dict(zip(names, seqs))}
            yield ("" if first else ",") + json.dumps(row, ensure_ascii=False)
            first = False
        yield '], "backup": {"segments": ' + json.dumps(store.segments, ensure_ascii=False)
        yield ', "monitors": ' + json.dumps(store.monitors, ensure_ascii=False) + ', "rounds": ['
        first = True
        for rid, t, code, src in store.iter_rounds(since):
            yield ("" if first else ",") + json.dumps({"id": rid, "t": int(float(t) * 1000), "code": code, "src": src})
            first = False
        yield "]}}"

    name = f"hugoland_{datetime.now(TZ).strftime('%Y-%m-%d_%H%M')}.json"
    return Response(gen(), mimetype="application/json", headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/api/import")
@need_login
def api_import():
    """Aceita o JSON exportado pela nuvem ou pelo app do tablet (versão 2)."""
    try:
        d = json.loads(request.get_data(as_text=True))
    except ValueError:
        return jsonify(error="arquivo não é JSON"), 400
    b = d.get("backup", d) if isinstance(d, dict) else {}
    rounds = b.get("rounds") or []
    old_segs = {s.get("id"): s for s in (b.get("segments") or [])}

    def code_of_old(seg_id):
        s = old_segs.get(seg_id)
        if not s:
            return None
        for a in str(s.get("aliases", "")).split(","):
            if re.fullmatch(r"ABW_[A-Z0-9_]+", a.strip()):
                return a.strip()
        cur = store.seg_for_code(s.get("label", ""))
        return store.code_for_seg(cur["id"]) if cur else None

    items_api, items_manual = [], []
    for r in rounds:
        try:
            t = float(r["t"]) / 1000
            code = r.get("code") or (code_of_old(r["v"][0]) if r.get("v") else None)
            if not code:
                continue
            rid = str(r.get("id") or f"imp-{int(t * 1000)}-{code}")
            (items_manual if r.get("src") == "manual" else items_api).append((rid, t, code))
        except (KeyError, TypeError, ValueError, IndexError):
            continue
    new = store.insert_rounds(items_api, "api") + store.insert_rounds(items_manual, "manual")
    # combinações: adiciona as que ainda não existem (pelo nome)
    names = {m["name"] for m in store.monitors}
    added = 0
    for m in b.get("monitors") or []:
        if m.get("name") in names:
            continue
        vals = []
        for v in m.get("values", []):
            if any(s["id"] == v for s in store.segments):
                vals.append(v)
            else:
                code = code_of_old(v)
                s = store.seg_for_code(code) if code else None
                if s:
                    vals.append(s["id"])
        if vals:
            store.upsert_monitor({"name": m["name"], "values": vals, "mode": m.get("mode", "any"), "alert": int(m.get("alert") or 0)})
            added += 1
    return jsonify(ok=True, rodadas=new, combinacoes=added)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")))
